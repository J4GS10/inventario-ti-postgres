"""
Cierre mensual para Contabilidad: el valor del inventario al cierre de un mes y
lo que cambió en el mes, por sede y por moneda (no se convierten monedas).

- Activos: saldo inicial (valor al día anterior al mes), altas del mes (por
  fecha de compra), bajas del mes (por fecha de baja) y saldo final; los que no
  tienen costo se cuentan aparte.
- Consumibles: por consumible, existencia y valor al inicio y al cierre (último
  saldo del kardex al costo promedio), entradas, salidas y ajustes del mes.
- Compras recibidas en el mes (recepciones al precio real o de la orden).
- Licencias vigentes al cierre con su costo, y mantenimientos cerrados en el mes.

Se exporta a Excel con una hoja por sección y el detalle de altas y bajas.
"""
from __future__ import annotations

import calendar
import io
from datetime import date, datetime, time
from decimal import Decimal

from sqlalchemy import and_, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.catalogs import EstadoOperativo, Modelo, TipoActivo
from app.models.consumable import Consumible
from app.models.core import Activo
from app.models.kardex import Kardex
from app.models.location import Sede
from app.models.procurement import OrdenCompra, OrdenCompraLinea, Proveedor, Recepcion, RecepcionLinea
from app.models.software import Licencia, Software
from app.models.traceability import Mantenimiento


def _d(v) -> Decimal:
    return Decimal(str(v or 0))


def _txt(v: Decimal) -> str:
    return str(v.quantize(Decimal("0.01")))


class CierreMensualService:
    def __init__(self, db: AsyncSession):
        self.db = db

    @staticmethod
    def periodo(anio: int, mes: int) -> tuple[date, date]:
        return date(anio, mes, 1), date(anio, mes, calendar.monthrange(anio, mes)[1])

    async def _sedes(self) -> dict[int, str]:
        return dict((await self.db.execute(select(Sede.SED_Sede, Sede.SED_Nombre))).all())

    # ------------------------------------------------------------------
    async def _activos(self, inicio: date, fin: date, sede: int | None) -> dict:
        def vigente_en(dia: date):
            return and_(Activo.ACT_Fecha_Compra <= dia,
                        or_(Activo.ACT_Baja_Fecha.is_(None), Activo.ACT_Baja_Fecha > dia))

        base = select(Activo.SED_Sede, Activo.ACT_Moneda, func.count(), func.coalesce(func.sum(Activo.ACT_Costo), 0),
                      func.count(Activo.ACT_Costo))
        if sede is not None:
            base = base.where(Activo.SED_Sede == sede)
        grupos: dict[tuple, dict] = {}

        async def sumar(clave: str, *cond):
            for s, moneda, n, valor, con_costo in (await self.db.execute(
                    base.where(*cond).group_by(Activo.SED_Sede, Activo.ACT_Moneda))).all():
                g = grupos.setdefault((s, moneda), {"SED_Sede": s, "moneda": moneda})
                g[f"{clave}_cantidad"] = g.get(f"{clave}_cantidad", 0) + n
                g[f"{clave}_valor"] = _d(g.get(f"{clave}_valor")) + _d(valor)
                g[f"{clave}_sin_costo"] = g.get(f"{clave}_sin_costo", 0) + (n - con_costo)

        dia_anterior = date.fromordinal(inicio.toordinal() - 1)
        await sumar("inicial", vigente_en(dia_anterior))
        await sumar("altas", Activo.ACT_Fecha_Compra.between(inicio, fin))
        await sumar("bajas", Activo.ACT_Baja_Fecha.between(inicio, fin), Activo.ACT_Fecha_Compra <= fin)
        await sumar("final", vigente_en(fin))
        resumen = []
        for g in grupos.values():
            fila = {"SED_Sede": g["SED_Sede"], "moneda": g["moneda"]}
            for clave in ("inicial", "altas", "bajas", "final"):
                fila[f"{clave}_cantidad"] = g.get(f"{clave}_cantidad", 0)
                fila[f"{clave}_valor"] = _txt(_d(g.get(f"{clave}_valor")))
                fila[f"{clave}_sin_costo"] = g.get(f"{clave}_sin_costo", 0)
            resumen.append(fila)

        async def detalle(*cond) -> list[dict]:
            q = (select(Activo, TipoActivo.TAC_Nombre, Modelo.MOD_Nombre)
                 .join(TipoActivo, TipoActivo.TAC_Tipo_Activo == Activo.TAC_Tipo_Activo)
                 .join(Modelo, Modelo.MOD_Modelo == Activo.MOD_Modelo).where(*cond).order_by(Activo.ACT_Codigo_Interno))
            if sede is not None:
                q = q.where(Activo.SED_Sede == sede)
            return [{"codigo": a.ACT_Codigo_Interno, "serie": a.ACT_Serie_Fabricante, "tipo": t, "modelo": m,
                     "SED_Sede": a.SED_Sede, "fecha_compra": a.ACT_Fecha_Compra, "fecha_baja": a.ACT_Baja_Fecha,
                     "motivo_baja": a.ACT_Baja_Motivo, "costo": _txt(_d(a.ACT_Costo)) if a.ACT_Costo is not None else None,
                     "moneda": a.ACT_Moneda, "factura": a.ACT_Factura}
                    for a, t, m in (await self.db.execute(q)).all()]

        return {
            "resumen": sorted(resumen, key=lambda r: (r["SED_Sede"] or 0, r["moneda"] or "")),
            "altas": await detalle(Activo.ACT_Fecha_Compra.between(inicio, fin)),
            "bajas": await detalle(Activo.ACT_Baja_Fecha.between(inicio, fin)),
        }

    # ------------------------------------------------------------------
    async def _consumibles(self, inicio: datetime, fin: datetime, sede: int | None) -> list[dict]:
        q = select(Consumible).order_by(Consumible.CON_Nombre)
        if sede is not None:
            q = q.where(Consumible.SED_Sede == sede)
        filas = []
        for c in (await self.db.execute(q)).scalars():
            async def saldo_a(momento: datetime):
                fila = (await self.db.execute(
                    select(Kardex.KDX_Saldo, Kardex.KDX_Saldo_Valor).where(
                        Kardex.CON_Consumible == c.CON_Consumible, Kardex.KDX_Fecha <= momento,
                        Kardex.KDX_Saldo.is_not(None))
                    .order_by(Kardex.KDX_Fecha.desc(), Kardex.KDX_Evento.desc()).limit(1))).first()
                return (fila[0], _d(fila[1])) if fila else (0, Decimal(0))

            ini_cant, ini_val = await saldo_a(inicio)
            fin_cant, fin_val = await saldo_a(fin)
            mov = {t: (0, Decimal(0)) for t in ("CONSUMIBLE_ENTRADA", "CONSUMIBLE_SALIDA", "CONSUMIBLE_AJUSTE")}
            for tipo, cant, valor in (await self.db.execute(
                    select(Kardex.KDX_Tipo, func.coalesce(func.sum(Kardex.KDX_Cantidad), 0),
                           func.coalesce(func.sum(func.abs(Kardex.KDX_Cantidad) * func.coalesce(Kardex.KDX_Costo, 0)), 0))
                    .where(Kardex.CON_Consumible == c.CON_Consumible, Kardex.KDX_Fecha > inicio, Kardex.KDX_Fecha <= fin,
                           Kardex.KDX_Tipo.in_(tuple(mov)))
                    .group_by(Kardex.KDX_Tipo))).all():
                mov[tipo] = (int(cant), _d(valor))
            if not any((ini_cant, fin_cant, *(m[0] for m in mov.values()))):
                continue
            filas.append({
                "CON_Consumible": c.CON_Consumible, "nombre": c.CON_Nombre, "unidad": c.CON_Unidad,
                "SED_Sede": c.SED_Sede, "moneda": c.CON_Moneda,
                "inicial_cantidad": ini_cant, "inicial_valor": _txt(ini_val),
                "entradas_cantidad": mov["CONSUMIBLE_ENTRADA"][0], "entradas_valor": _txt(mov["CONSUMIBLE_ENTRADA"][1]),
                "salidas_cantidad": mov["CONSUMIBLE_SALIDA"][0], "salidas_valor": _txt(mov["CONSUMIBLE_SALIDA"][1]),
                "ajustes_cantidad": mov["CONSUMIBLE_AJUSTE"][0],
                "final_cantidad": fin_cant, "final_valor": _txt(fin_val),
            })
        return filas

    # ------------------------------------------------------------------
    async def _compras(self, inicio: datetime, fin: datetime, sede: int | None) -> list[dict]:
        precio = func.coalesce(RecepcionLinea.RCL_Precio_Real, OrdenCompraLinea.OCL_Precio_Unitario)
        q = (select(OrdenCompra.OCO_Numero, Proveedor.PRV_Nombre, OrdenCompra.OCO_Moneda, OrdenCompra.SED_Sede,
                    func.max(Recepcion.REC_Fecha), func.sum(RecepcionLinea.RCL_Cantidad * precio))
             .join(Recepcion, Recepcion.OCO_Orden == OrdenCompra.OCO_Orden)
             .join(RecepcionLinea, RecepcionLinea.REC_Recepcion == Recepcion.REC_Recepcion)
             .join(OrdenCompraLinea, OrdenCompraLinea.OCL_Linea == RecepcionLinea.OCL_Linea)
             .join(Proveedor, Proveedor.PRV_Proveedor == OrdenCompra.PRV_Proveedor)
             .where(Recepcion.REC_Fecha > inicio, Recepcion.REC_Fecha <= fin)
             .group_by(OrdenCompra.OCO_Numero, Proveedor.PRV_Nombre, OrdenCompra.OCO_Moneda, OrdenCompra.SED_Sede)
             .order_by(OrdenCompra.OCO_Numero))
        if sede is not None:
            q = q.where(OrdenCompra.SED_Sede == sede)
        return [{"orden": n, "proveedor": p, "moneda": m, "SED_Sede": s, "fecha": f, "valor": _txt(_d(v))}
                for n, p, m, s, f, v in (await self.db.execute(q)).all()]

    async def _licencias(self, fin: date) -> list[dict]:
        q = (select(Software.SOF_Nombre, Licencia.LIC_Cantidad_Total, Licencia.LIC_Costo, Licencia.LIC_Moneda,
                    Licencia.LIC_Fecha_Vencimiento)
             .join(Software, Software.SOF_Software == Licencia.SOF_Software)
             .where(or_(Licencia.LIC_Fecha_Vencimiento.is_(None), Licencia.LIC_Fecha_Vencimiento >= fin))
             .order_by(Software.SOF_Nombre))
        return [{"software": s, "cantidad": n, "costo": _txt(_d(c)) if c is not None else None, "moneda": m,
                 "vencimiento": v} for s, n, c, m, v in (await self.db.execute(q)).all()]

    async def _mantenimiento(self, inicio: datetime, fin: datetime, sede: int | None) -> list[dict]:
        q = (select(Activo.SED_Sede, Mantenimiento.MAN_Moneda, func.count(), func.sum(Mantenimiento.MAN_Costo_Total))
             .join(Activo, Activo.ACT_Activo == Mantenimiento.ACT_Activo)
             .where(Mantenimiento.MAN_Fecha_Cierre > inicio, Mantenimiento.MAN_Fecha_Cierre <= fin)
             .group_by(Activo.SED_Sede, Mantenimiento.MAN_Moneda))
        if sede is not None:
            q = q.where(Activo.SED_Sede == sede)
        return [{"SED_Sede": s, "moneda": m, "tickets": n, "valor": _txt(_d(v))}
                for s, m, n, v in (await self.db.execute(q)).all()]

    # ------------------------------------------------------------------
    async def cierre(self, anio: int, mes: int, sede: int | None = None) -> dict:
        inicio, fin = self.periodo(anio, mes)
        inicio_dt = datetime.combine(date.fromordinal(inicio.toordinal() - 1), time.max)
        fin_dt = datetime.combine(fin, time.max)
        sedes = await self._sedes()
        activos = await self._activos(inicio, fin, sede)
        datos = {
            "anio": anio, "mes": mes, "desde": inicio, "hasta": fin, "SED_Sede": sede,
            "activos": activos["resumen"], "altas": activos["altas"], "bajas": activos["bajas"],
            "consumibles": await self._consumibles(inicio_dt, fin_dt, sede),
            "compras": await self._compras(inicio_dt, fin_dt, sede),
            "licencias": await self._licencias(fin),
            "mantenimiento": await self._mantenimiento(inicio_dt, fin_dt, sede),
        }
        for seccion in ("activos", "altas", "bajas", "consumibles", "compras", "mantenimiento"):
            for fila in datos[seccion]:
                fila["sede"] = sedes.get(fila.get("SED_Sede"))
        return datos

    async def cierre_xlsx(self, anio: int, mes: int, sede: int | None, lang: str) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from app.services.cierre_textos import TEXTOS
        t = TEXTOS[lang]
        datos = await self.cierre(anio, mes, sede)
        libro = Workbook()
        libro.remove(libro.active)

        def hoja(nombre: str, columnas: list[tuple[str, str]], filas: list[dict]) -> None:
            ws = libro.create_sheet(t[nombre][:31])
            ws.append([t[c] for _, c in columnas])
            for celda in ws[1]:
                celda.font = Font(bold=True)
            for fila in filas:
                ws.append([_celda(fila.get(k), numerico=k.endswith("valor") or k == "costo") for k, _ in columnas])
            for i, (_, c) in enumerate(columnas, start=1):
                ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, len(t[c]) + 2)

        hoja("hoja_activos", [("sede", "col_sede"), ("moneda", "col_moneda"), ("inicial_cantidad", "col_ini_n"),
                              ("inicial_valor", "col_ini_v"), ("altas_cantidad", "col_altas_n"),
                              ("altas_valor", "col_altas_v"), ("bajas_cantidad", "col_bajas_n"),
                              ("bajas_valor", "col_bajas_v"), ("final_cantidad", "col_fin_n"),
                              ("final_valor", "col_fin_v"), ("final_sin_costo", "col_sin_costo")], datos["activos"])
        detalle = [("codigo", "col_codigo"), ("serie", "col_serie"), ("tipo", "col_tipo"), ("modelo", "col_modelo"),
                   ("sede", "col_sede"), ("fecha_compra", "col_fecha_compra"), ("fecha_baja", "col_fecha_baja"),
                   ("motivo_baja", "col_motivo_baja"), ("costo", "col_costo"), ("moneda", "col_moneda"),
                   ("factura", "col_factura")]
        hoja("hoja_altas", detalle, datos["altas"])
        hoja("hoja_bajas", detalle, datos["bajas"])
        hoja("hoja_consumibles", [("nombre", "col_nombre"), ("sede", "col_sede"), ("moneda", "col_moneda"),
                                  ("inicial_cantidad", "col_ini_n"), ("inicial_valor", "col_ini_v"),
                                  ("entradas_cantidad", "col_ent_n"), ("entradas_valor", "col_ent_v"),
                                  ("salidas_cantidad", "col_sal_n"), ("salidas_valor", "col_sal_v"),
                                  ("ajustes_cantidad", "col_aju_n"), ("final_cantidad", "col_fin_n"),
                                  ("final_valor", "col_fin_v")], datos["consumibles"])
        hoja("hoja_compras", [("orden", "col_orden"), ("proveedor", "col_proveedor"), ("sede", "col_sede"),
                              ("fecha", "col_fecha"), ("moneda", "col_moneda"), ("valor", "col_valor")], datos["compras"])
        hoja("hoja_licencias", [("software", "col_software"), ("cantidad", "col_cantidad"), ("costo", "col_costo"),
                                ("moneda", "col_moneda"), ("vencimiento", "col_vencimiento")], datos["licencias"])
        hoja("hoja_mantenimiento", [("sede", "col_sede"), ("moneda", "col_moneda"), ("tickets", "col_tickets"),
                                    ("valor", "col_valor")], datos["mantenimiento"])
        salida = io.BytesIO()
        libro.save(salida)
        return salida.getvalue()


def _celda(v, numerico: bool = False):
    if isinstance(v, datetime):
        return v.replace(microsecond=0)
    if numerico and isinstance(v, str):
        return float(Decimal(v))
    return v
