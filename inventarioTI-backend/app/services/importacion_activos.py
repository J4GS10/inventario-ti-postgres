"""
Importación de activos desde Excel (.xlsx) con vista previa.

Flujo: el usuario descarga la plantilla (o usa su propia hoja con los mismos
encabezados, como la de «Inventario GT»), la sube, revisa fila por fila lo que
se va a crear y confirma. La confirmación vuelve a validar el archivo y crea
todo en una sola transacción usando el alta normal de activos (código
automático, kardex y auditoría incluidos).

Los activos entran disponibles en la ubicación indicada; las asignaciones se
hacen después con la entrega de puesto.
"""
from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from fastapi import HTTPException
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.currency import DEFAULT_CURRENCY, SUPPORTED_CURRENCIES as CURRENCY_CODES
from app.core.data_scope import area_sede_id, elevated
from app.core.transactional import transactional
from app.models.catalogs import EstadoOperativo, Marca, Modelo, TipoActivo, TipoEspecificacion
from app.models.core import Activo
from app.models.enums import EstadoOperativoEnum
from app.models.governance import Secuencia
from app.models.location import Area
from app.repositories.governance import GovernanceRepository

MAX_FILAS = 1000
MAX_BYTES = 5 * 1024 * 1024
SERIES_GENERICAS = {"", "generico", "n/a", "na", "s/n", "sin serie", "none", "-"}

# Encabezado normalizado -> campo. Acepta la plantilla y la hoja INVENTARIO del Excel original.
ENCABEZADOS = {
    "codigo": "codigo", "code": "codigo", "codice": "codigo",
    "categoria": "categoria", "tipo": "categoria", "tipo de activo": "categoria", "type": "categoria",
    "category": "categoria",
    "descripcion breve": "descripcion", "descripcion": "descripcion", "description": "descripcion",
    "marca": "marca", "brand": "marca",
    "modelo": "modelo", "model": "modelo",
    "serie": "serie", "numero de serie": "serie", "serie del fabricante": "serie", "serial": "serie",
    "hostname": "hostname", "host": "hostname",
    "ubicacion actual": "ubicacion", "ubicacion": "ubicacion", "area": "ubicacion", "location": "ubicacion",
    "fecha entrada": "fecha", "fecha de compra": "fecha", "fecha compra": "fecha", "fecha": "fecha",
    "purchase date": "fecha",
    "id factura": "factura", "factura": "factura", "invoice": "factura",
    "observaciones": "observaciones", "notas": "observaciones", "notes": "observaciones",
    "costo": "costo", "cost": "costo", "moneda": "moneda", "currency": "moneda",
}
COLUMNAS_PLANTILLA = ["Código", "Categoría", "Descripción", "Marca", "Modelo", "Serie", "Hostname",
                      "Ubicación", "Fecha de compra", "Factura", "Costo", "Moneda", "Observaciones"]


def _norm(v) -> str:
    s = unicodedata.normalize("NFKD", str(v if v is not None else ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().casefold()


def _clave_codigo(codigo: str) -> str:
    """Forma comparable de un código: «CAM-0001» y «cam-00001» son el mismo."""
    c = re.sub(r"\s+", "", codigo or "").upper()
    m = re.fullmatch(r"([A-Z]+-?)0*(\d+)", c)
    return f"{m.group(1)}{int(m.group(2))}" if m else c


def _texto(v) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


@dataclass
class Fila:
    fila: int
    datos: dict
    errores: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)
    # Resueltos durante la validación
    tipo: TipoActivo | None = None
    area: Area | None = None
    marca_nombre: str = ""
    modelo_nombre: str = ""
    fecha: date | None = None
    costo: Decimal | None = None
    moneda: str = DEFAULT_CURRENCY
    marca_nueva: str | None = None
    modelo_nuevo: tuple[str, str] | None = None

    def resumen(self) -> dict:
        return {
            "fila": self.fila,
            "codigo": self.datos.get("codigo") or None,
            "categoria": self.tipo.TAC_Nombre if self.tipo else self.datos.get("categoria") or None,
            "marca": self.marca_nombre or None,
            "modelo": self.modelo_nombre or None,
            "serie": self.datos.get("serie") or None,
            "ubicacion": self.area.ARE_Nombre if self.area else self.datos.get("ubicacion") or None,
            "estado": "error" if self.errores else ("aviso" if self.avisos else "ok"),
            "errores": self.errores,
            "avisos": self.avisos,
        }


def leer_excel(contenido: bytes) -> list[Fila]:
    """Encuentra la hoja y la fila de encabezados y devuelve las filas con datos."""
    if len(contenido) > MAX_BYTES:
        raise HTTPException(413, detail="FILE_TOO_LARGE")
    try:
        import openpyxl
        libro = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - archivo dañado o no es .xlsx
        raise HTTPException(400, detail="IMPORT_INVALID_FILE") from exc

    hojas = sorted(libro.worksheets, key=lambda h: 0 if _norm(h.title) == "inventario" else 1)
    for hoja in hojas:
        filas = list(hoja.iter_rows(max_row=MAX_FILAS + 20, values_only=True))
        for i, fila in enumerate(filas[:15]):
            mapa = {}
            for j, celda in enumerate(fila):
                campo = ENCABEZADOS.get(_norm(celda))
                if campo and campo not in mapa:
                    mapa[campo] = j
            if {"categoria", "marca"} <= mapa.keys() and ({"modelo", "descripcion"} & mapa.keys()):
                resultado = []
                for k, datos in enumerate(filas[i + 1:], start=i + 2):
                    valores = {c: datos[j] if j < len(datos) else None for c, j in mapa.items()}
                    if not any(_texto(v) for v in valores.values()):
                        continue
                    resultado.append(Fila(fila=k, datos={
                        c: (v if c in ("fecha", "costo") else _texto(v)) for c, v in valores.items()
                    }))
                if len(resultado) > MAX_FILAS:
                    raise HTTPException(400, detail=f"IMPORT_TOO_MANY_ROWS:{MAX_FILAS}")
                if not resultado:
                    raise HTTPException(400, detail="IMPORT_EMPTY")
                return resultado
    raise HTTPException(400, detail="IMPORT_HEADERS_NOT_FOUND")


class ImportacionActivosService:
    def __init__(self, db: AsyncSession):
        self.db = db
        self.gov_repo = GovernanceRepository(db)

    async def _catalogos(self):
        tipos = (await self.db.execute(select(TipoActivo))).scalars().all()
        marcas = (await self.db.execute(select(Marca))).scalars().all()
        modelos = (await self.db.execute(select(Modelo))).scalars().all()
        areas = (await self.db.execute(select(Area))).scalars().all()
        return (
            {_norm(t.TAC_Nombre): t for t in tipos},
            {_norm(m.MAR_Nombre): m for m in marcas},
            {(m.MAR_Marca, _norm(m.MOD_Nombre)): m for m in modelos},
            {_norm(a.ARE_Nombre): a for a in areas},
        )

    async def analizar(self, contenido: bytes) -> tuple[list[Fila], dict]:
        filas = leer_excel(contenido)
        tipos, marcas, modelos, areas = await self._catalogos()
        series = [f.datos.get("serie", "") for f in filas if _norm(f.datos.get("serie")) not in SERIES_GENERICAS]
        async with elevated(self.db):
            # Unicidad en toda la empresa y sin importar ceros a la izquierda.
            codigos_db = {_clave_codigo(c) for c in (await self.db.execute(select(Activo.ACT_Codigo_Interno))).scalars().all()}
            series_db = set((await self.db.execute(
                select(func.lower(Activo.ACT_Serie_Fabricante))
                .where(func.lower(Activo.ACT_Serie_Fabricante).in_([s.lower() for s in series]))
            )).scalars().all()) if series else set()

        vistos_cod: set[str] = set()
        vistas_serie: set[str] = set()
        for f in filas:
            d = f.datos
            d["codigo"] = re.sub(r"\s+", "", d.get("codigo") or "").upper()
            # Tipo
            if not d.get("categoria"):
                f.errores.append("MISSING_TYPE")
            else:
                f.tipo = tipos.get(_norm(d["categoria"]))
                if f.tipo is None and d.get("codigo"):
                    # El prefijo del código identifica el tipo aunque el nombre difiera
                    # («Power Adapter» en el Excel, «Cargador de laptop» en el sistema).
                    codigo_mayus = d["codigo"].upper()
                    candidatos = [t for t in tipos.values()
                                  if t.TAC_Prefijo and codigo_mayus.startswith(t.TAC_Prefijo.upper())]
                    if candidatos:
                        f.tipo = max(candidatos, key=lambda t: len(t.TAC_Prefijo))
                        f.avisos.append("TYPE_FROM_CODE")
                if f.tipo is None:
                    f.errores.append("TYPE_NOT_FOUND")
            # Marca y modelo (se crean si no existen)
            f.marca_nombre = d.get("marca") or ""
            if not f.marca_nombre:
                f.errores.append("MISSING_BRAND")
            desc, cod = d.get("descripcion") or "", d.get("modelo") or ""
            if _norm(cod) in SERIES_GENERICAS:
                cod = ""
            f.modelo_nombre = (f"{desc} ({cod})" if desc and cod and _norm(cod) not in _norm(desc)
                               else (desc or cod))[:150]
            if not f.modelo_nombre:
                f.errores.append("MISSING_MODEL")
            if f.marca_nombre and f.modelo_nombre:
                marca = marcas.get(_norm(f.marca_nombre))
                if marca is None:
                    f.marca_nueva = _norm(f.marca_nombre)
                    f.avisos.append("NEW_BRAND")
                if marca is None or (marca.MAR_Marca, _norm(f.modelo_nombre)) not in modelos:
                    f.modelo_nuevo = (_norm(f.marca_nombre), _norm(f.modelo_nombre))
                    f.avisos.append("NEW_MODEL")
            # Código
            codigo = d["codigo"]
            if codigo:
                if _clave_codigo(codigo) in codigos_db:
                    f.errores.append("CODE_EXISTS")
                elif _clave_codigo(codigo) in vistos_cod:
                    f.errores.append("CODE_DUPLICATED_IN_FILE")
                vistos_cod.add(_clave_codigo(codigo))
                if f.tipo and f.tipo.TAC_Prefijo and not codigo.startswith(f.tipo.TAC_Prefijo.upper()):
                    f.avisos.append("CODE_PREFIX_MISMATCH")
            elif f.tipo and not f.tipo.TAC_Prefijo:
                f.errores.append("TYPE_WITHOUT_PREFIX")
            else:
                f.avisos.append("AUTO_CODE")
            # Serie
            serie = d.get("serie") or ""
            if _norm(serie) in SERIES_GENERICAS:
                d["serie"] = ""
                f.avisos.append("NO_SERIAL_USES_CODE")
            elif serie.lower() in series_db:
                f.errores.append("SERIAL_EXISTS")
            elif serie.lower() in vistas_serie:
                f.errores.append("SERIAL_DUPLICATED_IN_FILE")
            elif len(serie) < 3:
                f.errores.append("SERIAL_TOO_SHORT")
            vistas_serie.add(serie.lower())
            # Ubicación
            if d.get("ubicacion"):
                f.area = areas.get(_norm(d["ubicacion"]))
                if f.area is None:
                    f.errores.append("AREA_NOT_FOUND")
            else:
                f.errores.append("MISSING_AREA")
            # Fecha
            valor = d.get("fecha")
            if isinstance(valor, datetime):
                f.fecha = valor.date()
            elif isinstance(valor, date):
                f.fecha = valor
            elif _texto(valor):
                try:
                    f.fecha = datetime.strptime(_texto(valor)[:10], "%Y-%m-%d").date()
                except ValueError:
                    try:
                        f.fecha = datetime.strptime(_texto(valor)[:10], "%d/%m/%Y").date()
                    except ValueError:
                        f.errores.append("INVALID_DATE")
            else:
                f.fecha = date.today()
                f.avisos.append("NO_DATE_USES_TODAY")
            # Costo y moneda
            if _texto(d.get("costo")):
                try:
                    f.costo = Decimal(str(d["costo"]).replace(",", "")).quantize(Decimal("0.01"))
                    if f.costo < 0:
                        raise InvalidOperation
                except (InvalidOperation, ValueError):
                    f.errores.append("INVALID_COST")
            moneda = (d.get("moneda") or "").upper()
            if moneda and moneda not in CURRENCY_CODES:
                f.errores.append("INVALID_CURRENCY")
            f.moneda = moneda or DEFAULT_CURRENCY

        resumen = {
            "filas": len(filas),
            "validas": sum(1 for f in filas if not f.errores),
            "con_error": sum(1 for f in filas if f.errores),
            # Solo lo que realmente se creará (filas sin error).
            "marcas_nuevas": len({f.marca_nueva for f in filas if f.marca_nueva and not f.errores}),
            "modelos_nuevos": len({f.modelo_nuevo for f in filas if f.modelo_nuevo and not f.errores}),
        }
        return filas, resumen

    async def previsualizar(self, contenido: bytes) -> dict:
        filas, resumen = await self.analizar(contenido)
        return {"resumen": resumen, "filas": [f.resumen() for f in filas]}

    @transactional
    async def confirmar(self, contenido: bytes, *, omitir_errores: bool, usuario_id=None, ip=None) -> dict:
        from app.schemas.core import ActivoCreate
        from app.schemas.core import EspecificacionCreate
        from app.services.core import CoreService

        filas, resumen = await self.analizar(contenido)
        if resumen["con_error"] and not omitir_errores:
            raise HTTPException(400, detail="IMPORT_HAS_ERRORS")
        validas = [f for f in filas if not f.errores]
        if not validas:
            raise HTTPException(400, detail="IMPORT_NOTHING_TO_IMPORT")

        disponible = (await self.db.execute(
            select(EstadoOperativo.EOP_Estado_Operativo)
            .where(EstadoOperativo.EOP_Nombre == EstadoOperativoEnum.DISPONIBLE.value)
        )).scalar_one()
        tes = {}
        for nombre in ("Factura", "Observaciones", "Nota de serie"):
            obj = (await self.db.execute(
                select(TipoEspecificacion).where(func.lower(TipoEspecificacion.TES_Nombre) == nombre.lower())
            )).scalar_one_or_none()
            if obj is None:
                obj = TipoEspecificacion(TES_Nombre=nombre)
                self.db.add(obj)
                await self.db.flush()
            tes[nombre] = obj.TES_Tipo_Especificacion

        _, marcas, modelos, _ = await self._catalogos()
        core = CoreService(self.db)
        creados: list[str] = []
        for f in validas:
            marca = marcas.get(_norm(f.marca_nombre))
            if marca is None:
                marca = Marca(MAR_Nombre=f.marca_nombre[:100])
                self.db.add(marca)
                await self.db.flush()
                marcas[_norm(f.marca_nombre)] = marca
            modelo = modelos.get((marca.MAR_Marca, _norm(f.modelo_nombre)))
            if modelo is None:
                modelo = Modelo(MOD_Nombre=f.modelo_nombre, MAR_Marca=marca.MAR_Marca,
                                TAC_Tipo_Activo=f.tipo.TAC_Tipo_Activo)
                self.db.add(modelo)
                await self.db.flush()
                modelos[(marca.MAR_Marca, _norm(f.modelo_nombre))] = modelo

            codigo = f.datos.get("codigo") or await self.gov_repo.get_next_code(
                contexto=f"ASSET_{f.tipo.TAC_Prefijo}", prefijo=f.tipo.TAC_Prefijo)
            serie = f.datos.get("serie") or codigo
            specs = [EspecificacionCreate(TES_Tipo_Especificacion=tes[n], ESP_Valor=v[:255]) for n, v in (
                ("Factura", f.datos.get("factura")), ("Observaciones", f.datos.get("observaciones")),
                ("Nota de serie", None if f.datos.get("serie") else "Sin serie del fabricante: se usa el código interno."),
            ) if v]
            sede = await area_sede_id(self.db, f.area.ARE_Area)
            try:
                activo = await core.create_activo(ActivoCreate(
                    ACT_Codigo_Interno=codigo, ACT_Serie_Fabricante=serie,
                    ACT_Hostname=f.datos.get("hostname") or None, ACT_Fecha_Compra=f.fecha,
                    ACT_Costo=f.costo, ACT_Moneda=f.moneda, MOD_Modelo=modelo.MOD_Modelo,
                    TAC_Tipo_Activo=f.tipo.TAC_Tipo_Activo, EOP_Estado_Operativo=disponible,
                    SED_Sede=sede, especificaciones=specs,
                ), usuario_id=usuario_id)
            except HTTPException as exc:
                raise HTTPException(exc.status_code, detail=f"IMPORT_ROW_FAILED:{f.fila}:{exc.detail}") from exc
            activo.ARE_Area_Ubicacion = f.area.ARE_Area
            creados.append(codigo)
        await self.db.flush()
        await self._ajustar_secuencias()
        await self.gov_repo.create_audit_log(
            "ASSET_IMPORT", "INV_ACTIVO",
            {"creados": len(creados), "omitidos": len(filas) - len(validas), "codigos": creados[:200]},
            usuario_id=usuario_id, ip_origen=ip,
        )
        return {"creados": len(creados), "omitidos": len(filas) - len(validas), "codigos": creados}

    async def _ajustar_secuencias(self) -> None:
        """Lleva cada secuencia al mayor número usado, así el próximo código automático no choca."""
        tipos = (await self.db.execute(select(TipoActivo).where(TipoActivo.TAC_Prefijo.is_not(None)))).scalars().all()
        for t in tipos:
            async with elevated(self.db):
                codigos = (await self.db.execute(
                    select(Activo.ACT_Codigo_Interno).where(Activo.ACT_Codigo_Interno.ilike(f"{t.TAC_Prefijo}%"))
                )).scalars().all()
            numeros = [int(m.group(1)) for c in codigos
                       if (m := re.fullmatch(re.escape(t.TAC_Prefijo) + r"(\d+)", c, flags=re.I))]
            if not numeros:
                continue
            contexto = f"ASSET_{t.TAC_Prefijo}"
            sec = (await self.db.execute(select(Secuencia).where(Secuencia.SEC_Contexto == contexto))).scalar_one_or_none()
            if sec is None:
                self.db.add(Secuencia(SEC_Contexto=contexto, SEC_Ultimo_Numero=max(numeros), SEC_Relleno=5))
            elif sec.SEC_Ultimo_Numero < max(numeros):
                sec.SEC_Ultimo_Numero = max(numeros)
        await self.db.flush()

    async def plantilla(self) -> bytes:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        tipos = (await self.db.execute(select(TipoActivo.TAC_Nombre).order_by(TipoActivo.TAC_Nombre))).scalars().all()
        areas = (await self.db.execute(select(Area.ARE_Nombre).order_by(Area.ARE_Nombre))).scalars().all()
        marcas = (await self.db.execute(select(Marca.MAR_Nombre).order_by(Marca.MAR_Nombre))).scalars().all()
        libro = openpyxl.Workbook()
        hoja = libro.active
        hoja.title = "INVENTARIO"
        hoja.append(COLUMNAS_PLANTILLA)
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="3F4E5F")
        for i, ancho in enumerate([12, 18, 36, 16, 18, 22, 14, 18, 16, 14, 10, 9, 30], start=1):
            hoja.column_dimensions[openpyxl.utils.get_column_letter(i)].width = ancho
        listas = libro.create_sheet("Listas")
        listas.append(["Categorías", "Ubicaciones", "Marcas existentes", "Monedas"])
        for i in range(max(len(tipos), len(areas), len(marcas), len(CURRENCY_CODES))):
            listas.append([
                tipos[i] if i < len(tipos) else None, areas[i] if i < len(areas) else None,
                marcas[i] if i < len(marcas) else None, CURRENCY_CODES[i] if i < len(CURRENCY_CODES) else None,
            ])
        salida = io.BytesIO()
        libro.save(salida)
        return salida.getvalue()
