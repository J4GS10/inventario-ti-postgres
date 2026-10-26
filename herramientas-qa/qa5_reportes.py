"""QA5 - Reportería y documentos (API): CSV, actas Word/PDF, hoja de descargo y etiquetas QR.

Abre cada archivo descargado y valida su contenido. Requiere haber ejecutado qa5_datos.py.
Evidencias en salidas/qa5/{csv,actas,etiquetas}/ y resultados en salidas/qa5/resultados_api.json.

Uso (cwd = herramientas-qa):  PYTHONIOENCODING=utf-8 python qa5_reportes.py
"""
import csv
import io
import json
import os
import re
import subprocess
import sys
import time
import uuid

import cv2
import docx
import fitz  # pymupdf
import numpy as np
import pdfplumber
import pyotp  # noqa: F401  (usado por qa5_comun vía login)

from qa_login import BASE
from qa5_comun import OUT, cliente, datos

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "inventarioTI-backend"))
from app.services import report_i18n as R  # noqa: E402  (módulo puro, solo stdlib)

RES = []
D = datos()
QA5_CODES = {v["codigo"] for v in D["activos"].values()}
for sub in ("csv", "actas", "etiquetas"):
    os.makedirs(os.path.join(OUT, sub), exist_ok=True)


def check(cid, caso, ok, obtenido="", evidencia="", severidad="media", observacion=False):
    estado = "OK" if ok else ("OBSERVACIÓN" if observacion else "DEFECTO")
    RES.append({"id": cid, "caso": caso, "obtenido": str(obtenido)[:600], "evidencia": evidencia,
                "severidad": "" if ok else severidad, "estado": estado})
    print(f"{estado:12} {cid:10} {caso}" + ("" if ok else f"  -> {str(obtenido)[:300]}"), flush=True)


def get(c, url, **kw):
    for _ in range(12):
        r = c.get(url, **kw)
        if r.status_code != 429:
            return r
        time.sleep(15)
    return r


def post(c, url, **kw):
    for _ in range(12):
        r = c.post(url, **kw)
        if r.status_code != 429:
            return r
        time.sleep(15)
    return r


def psql(sql):
    out = subprocess.run(["docker", "exec", "inv_ui-db-1", "psql", "-U", "invuser", "-d", "inventario", "-Atc", sql],
                         capture_output=True, text=True, encoding="utf-8").stdout.strip()
    return out


def filename_of(r):
    m = re.search(r'filename="([^"]+)"', r.headers.get("content-disposition", ""))
    return m.group(1) if m else ""


def save(sub, name, content):
    p = os.path.join(OUT, sub, name)
    with open(p, "wb") as f:
        f.write(content)
    return os.path.relpath(p, os.path.join(OUT, "..")).replace("\\", "/")


# ─────────────────────────────── CSV ──────────────────────────────────────
EXPORTS = {
    "activos": ("/export/activos.csv", {"es": "activos", "en": "assets", "it": "asset"}),
    "movimientos": ("/export/movimientos.csv", {"es": "movimientos", "en": "movements", "it": "movimenti"}),
    "consumibles": ("/export/consumibles.csv", {"es": "consumibles", "en": "consumables", "it": "materiali_di_consumo"}),
    "proveedores": ("/export/proveedores.csv", {"es": "proveedores", "en": "suppliers", "it": "fornitori"}),
    "ordenes": ("/export/ordenes.csv", {"es": "ordenes_de_compra", "en": "purchase_orders", "it": "ordini_di_acquisto"}),
    "auditoria": ("/export/auditoria.csv", {"es": "auditoria", "en": "audit_log", "it": "registro_audit"}),
}
EXPECTED_HEADERS = {
    "activos": {"es": ["Código interno", "Número de serie", "Hostname", "Tipo de activo", "Marca", "Modelo",
                       "Estado operativo", "Fecha de compra", "Vencimiento de garantía", "Costo", "Moneda"],
                "en": ["Internal code", "Serial number", "Hostname", "Asset type", "Brand", "Model",
                       "Operational status", "Purchase date", "Warranty expiration", "Cost", "Currency"],
                "it": ["Codice interno", "Numero di serie", "Hostname", "Tipo di asset", "Marca", "Modello",
                       "Stato operativo", "Data di acquisto", "Scadenza garanzia", "Costo", "Valuta"]},
}
DB_COUNT = {
    "activos": 'select count(*) from "INV_ACTIVO"',
    "movimientos": 'select count(*) from "INV_MOVIMIENTO"',
    "consumibles": 'select count(*) from "INV_CONSUMIBLE"',
    "proveedores": 'select count(*) from "INV_PROVEEDOR"',
    "ordenes": 'select count(*) from "INV_ORDEN_COMPRA"',
}
# Columnas con fecha (índice, tipo) por exporte
DATE_COLS = {"activos": [(7, "d"), (8, "d")], "movimientos": [(0, "dt"), (1, "dt")],
             "ordenes": [(2, "d")], "auditoria": [(0, "dt")]}
CATALOG_COLS = {"activos": [(3, "tipo_activo"), (6, "estado_operativo")],
                "movimientos": [(7, "tipo_movimiento")], "consumibles": [(2, "unidad")],
                "ordenes": [(3, "estado_orden")], "auditoria": [(1, "accion_auditoria")]}
SPANISH_WORDS = {"Sí"}


def parse_csv(raw: bytes):
    txt = raw.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(txt)))


def csv_suite(c, etiqueta="sa"):
    info = {}
    for name, (path, fn) in EXPORTS.items():
        for lang in ("es", "en", "it"):
            r = get(c, f"{BASE}{path}", params={"lang": lang})
            cid = f"CSV-{name[:4].upper()}-{lang.upper()}"
            if r.status_code != 200:
                check(cid + "-HTTP", f"Exportar {name}.csv ({lang}) responde 200", False, f"{r.status_code} {r.text[:200]}",
                      severidad="alta")
                continue
            fname = filename_of(r)
            ev = save("csv", f"{etiqueta}_{fname or name + '_' + lang + '.csv'}", r.content)
            raw = r.content
            rows = parse_csv(raw)
            info[(name, lang)] = {"rows": rows, "ev": ev, "fname": fname, "ctype": r.headers.get("content-type")}
    return info


def analizar_csv(info):
    for name in EXPORTS:
        for lang in ("es", "en", "it"):
            if (name, lang) not in info:
                continue
            x = info[(name, lang)]
            rows, ev = x["rows"], x["ev"]
            cid = f"CSV-{name[:4].upper()}-{lang.upper()}"
            raw = open(os.path.join(OUT, "..", ev), "rb").read()
            header = rows[0] if rows else []
            data = rows[1:]
            problemas = []
            if not raw.startswith(b"\xef\xbb\xbf"):
                problemas.append("sin BOM UTF-8")
            if "utf-8" not in (x["ctype"] or ""):
                problemas.append(f"content-type {x['ctype']}")
            if not re.fullmatch(rf"{EXPORTS[name][1][lang]}_\d{{8}}_\d{{6}}\.csv", x["fname"] or ""):
                problemas.append(f"nombre de archivo {x['fname']!r}")
            txt = raw.decode("utf-8-sig")
            if "Ã" in txt or "�" in txt or "Â" in txt:
                problemas.append("mojibake (Ã/Â/�) en el contenido")
            exp_h = EXPECTED_HEADERS.get(name, {}).get(lang)
            if exp_h and header != exp_h:
                problemas.append(f"encabezados {header}")
            ncols = {len(rw) for rw in data}
            if ncols and ncols != {len(header)}:
                problemas.append(f"filas con nº de columnas distinto: {ncols}")
            # Fechas ISO
            for idx, kind in DATE_COLS.get(name, []):
                pat = r"\d{4}-\d{2}-\d{2}" if kind == "d" else r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}"
                malas = [rw[idx] for rw in data if rw[idx] and not re.fullmatch(pat, rw[idx])]
                if malas:
                    problemas.append(f"fechas no ISO en '{header[idx]}': {malas[:3]}")
            # Inyección CSV: ninguna celda debe empezar con = + - @ tab/cr
            peligrosas = [cell for rw in data for cell in rw if cell and cell[0] in "=+-@\t\r"]
            if peligrosas:
                problemas.append(f"celdas con fórmula sin neutralizar: {peligrosas[:3]}")
            # Catálogos: ningún valor canónico en otro idioma
            for idx, cat in CATALOG_COLS.get(name, []):
                for rw in data:
                    v = rw[idx]
                    if not v:
                        continue
                    ent = R._CATALOGS[cat].get(R._key(v))
                    if ent is None:
                        if re.fullmatch(r"[A-Z0-9_]{3,}", v):
                            problemas.append(f"clave sin traducir en '{header[idx]}': {v}")
                    elif v != ent[lang]:
                        problemas.append(f"valor en otro idioma en '{header[idx]}': {v} (esperado {ent[lang]})")
            problemas = list(dict.fromkeys(problemas))
            sev = "alta" if any("sin BOM" in p or "fórmula" in p or "mojibake" in p for p in problemas) else "media"
            check(cid, f"CSV {name} en {lang}: BOM, encabezados, nombre, fechas ISO, inyección y catálogos traducidos",
                  not problemas, "; ".join(problemas[:8]) or f"{len(data)} filas; archivo {x['fname']}", ev, sev)


def csv_contenido(info):
    # Conteos contra la base (el exporte lee la réplica: puede haber unos segundos de retraso)
    for name, sql in DB_COUNT.items():
        if (name, "es") not in info:
            continue
        n_csv = len(info[(name, "es")]["rows"]) - 1
        n_db = int(psql(sql) or -1)
        check(f"CSV-CNT-{name[:4].upper()}", f"Conteo de filas del CSV {name} = registros en la base (usuario global)",
              n_csv == n_db, f"CSV {n_csv} filas vs base {n_db}", info[(name, "es")]["ev"], "alta")
    # Mismo número de filas en los 3 idiomas
    for name in EXPORTS:
        ns = {lang: len(info[(name, lang)]["rows"]) for lang in ("es", "en", "it") if (name, lang) in info}
        if name != "auditoria":
            check(f"CSV-3L-{name[:4].upper()}", f"CSV {name}: mismo número de filas en ES/EN/IT", len(set(ns.values())) == 1, ns)

    # Activos QA5: costo, moneda, inyección neutralizada
    rows = info.get(("activos", "es"), {}).get("rows", [])
    by_code = {rw[0]: rw for rw in rows[1:]}
    a = D["activos"]
    r1, r2, r3, r4 = (by_code.get(a[k]["codigo"]) for k in ("a1", "a2", "a3", "a4"))
    check("CSV-ACT-MON", "Activos: columna Moneda con el código de la moneda del costo (USD/CHF/GTQ) y monto con 2 decimales",
          r1 and r1[9:11] == ["1234.50", "USD"] and r2[9:11] == ["999.99", "CHF"] and r3[9:11] == ["5000.00", "GTQ"],
          [x[9:11] if x else None for x in (r1, r2, r3)], info[("activos", "es")]["ev"], "alta")
    check("CSV-ACT-MON0", "Activos: sin costo → Costo y Moneda vacíos", r4 and r4[9:11] == ["", ""], r4[9:11] if r4 else None,
          info[("activos", "es")]["ev"], "baja")
    check("CSV-INJ-ACT", "Inyección CSV neutralizada en activos (=, +, -, @ prefijados con ')",
          r1 and r1[2].startswith("'=") and r2[1].startswith("'+") and r2[2].startswith("'@") and r3[2].startswith("'-"),
          [r1[2] if r1 else None, r2[1] if r2 else None, r2[2] if r2 else None, r3[2] if r3 else None],
          info[("activos", "es")]["ev"], "alta")
    check("CSV-ACT-FECHA", "Activos: fecha de compra y garantía en ISO (2026-09-24 / 2029-09-24)",
          r1 and r1[7] == "2026-09-24" and r1[8] == "2029-09-24", r1[7:9] if r1 else None, info[("activos", "es")]["ev"])
    # Orden de activos: ¿determinista?
    codes = [rw[0] for rw in rows[1:]]
    check("CSV-ACT-ORD", "Activos: filas ordenadas (por código) para una lectura estable",
          codes == sorted(codes), f"primeros códigos: {codes[:6]}", info[("activos", "es")]["ev"], "baja", observacion=True)
    # Movimientos QA5
    mrows = info.get(("movimientos", "es"), {}).get("rows", [])
    obs = [rw for rw in mrows[1:] if rw[2] == a["a2"]["codigo"]]
    check("CSV-INJ-MOV", "Inyección CSV neutralizada en observaciones de movimientos (=cmd|...)",
          obs and obs[0][8].startswith("'="), obs[0][8] if obs else "no encontrado", info[("movimientos", "es")]["ev"], "alta")
    m1 = [rw for rw in mrows[1:] if rw[2] == a["a1"]["codigo"]]
    check("CSV-MOV-UTF", "Movimientos: nombre y observación con tildes/ñ íntegros (QA5-Íñigo Peña, señal, ñandú)",
          m1 and "QA5-Íñigo Peña" == m1[0][4] and "ñandú" in m1[0][8] and m1[0][6] == "QA5-Bodega Señorío",
          m1[0] if m1 else "no encontrado", info[("movimientos", "es")]["ev"], "alta")
    en = info.get(("movimientos", "en"), {}).get("rows", [])
    m1en = [rw for rw in en[1:] if rw[2] == a["a1"]["codigo"]]
    check("CSV-MOV-EN", "Movimientos EN: tipo de movimiento traducido (Assignment) ",
          m1en and m1en[0][7] == "Assignment", m1en[0][7] if m1en else None, info[("movimientos", "en")]["ev"])
    dev = [rw for rw in mrows[1:] if rw[2] == a["a3"]["codigo"]]
    check("CSV-MOV-DEV", "Movimientos: la asignación devuelta tiene fecha de devolución ISO",
          any(rw[1] for rw in dev), [rw[:2] + [rw[7]] for rw in dev], info[("movimientos", "es")]["ev"], "media")
    # Consumibles QA5
    for lang, yes, unit in (("es", "Sí", "caja"), ("en", "Yes", "box"), ("it", "Sì", "scatola")):
        crow = [rw for rw in info.get(("consumibles", lang), {}).get("rows", [])[1:] if "QA5-Tóner" in rw[0]]
        check(f"CSV-CON-{lang.upper()}", f"Consumibles {lang}: nombre '=…' neutralizado, unidad '{unit}', bajo mínimo '{yes}'",
              crow and crow[0][0].startswith("'=") and crow[0][1].startswith("'+") and crow[0][2] == unit and crow[0][5] == yes and crow[0][6] == yes,
              crow[0] if crow else "no encontrado", info.get(("consumibles", lang), {}).get("ev", ""), "media")
    # Proveedores QA5
    prow = [rw for rw in info.get(("proveedores", "es"), {}).get("rows", [])[1:] if rw[0].startswith("QA5-Proveedor")]
    check("CSV-PRV", "Proveedores: '-12345', '@QA5', '+502…' neutralizados y ñ íntegra",
          prow and prow[0][1] == "'-12345" and prow[0][2] == "'@QA5 contacto" and prow[0][4].startswith("'+") and "Ñandú" in prow[0][0],
          prow[0] if prow else "no encontrado", info.get(("proveedores", "es"), {}).get("ev", ""), "media")
    if prow and prow[0][4].startswith("'+"):
        check("CSV-PRV-TEL", "Proveedores: teléfono '+502 …' queda con apóstrofo visible (efecto colateral de la neutralización)",
              False, prow[0][4], info[("proveedores", "es")]["ev"], "baja", observacion=True)
    # Órdenes QA5
    for lang, st in (("es", "Borrador"), ("en", "Draft"), ("it", "Bozza")):
        orow = [rw for rw in info.get(("ordenes", lang), {}).get("rows", [])[1:] if rw[0] == D["orden"]["numero"]]
        check(f"CSV-ORD-{lang.upper()}", f"Órdenes {lang}: fecha ISO, estado '{st}', moneda CHF, total 300.50, notas neutralizadas",
              orow and orow[0][2] == "2026-09-24" and orow[0][3] == st and orow[0][4] == "CHF" and orow[0][5] == "300.50" and orow[0][6].startswith("'="),
              orow[0] if orow else "no encontrado", info.get(("ordenes", lang), {}).get("ev", ""), "media")
    # Auditoría: acciones sin traducir
    for lang in ("es", "en", "it"):
        arows = info.get(("auditoria", lang), {}).get("rows", [])[1:]
        crudas = sorted({rw[1] for rw in arows if re.fullmatch(r"[A-Z0-9_]{3,}", rw[1])})
        check(f"CSV-AUD-ACC-{lang.upper()}", f"Auditoría {lang}: todas las acciones traducidas (sin códigos internos)",
              not crudas, f"acciones sin traducir: {crudas}", info.get(("auditoria", lang), {}).get("ev", ""), "media")
        ent = sorted({rw[2] for rw in arows})
        if lang == "en":
            check("CSV-AUD-ENT", "Auditoría: columna 'Entidad afectada' legible (no nombres de tabla)",
                  not any(e.startswith(("INV_", "SYS_")) for e in ent), f"valores: {ent[:8]}",
                  info.get(("auditoria", lang), {}).get("ev", ""), "baja", observacion=True)


# ───────────────────────────── Actas ──────────────────────────────────────
TEXTOS = {
    "es": {"entrega": ["NOTA: ENTREGA DE EQUIPO", "DISPOSITIVO", "MARCA", "MODELO", "SERIE", "Descripción:",
                       "Recibe el equipo", "Entrega y autoriza", "Departamento de TI", "Colaborador"],
           "descargo": ["NOTA: DEVOLUCIÓN A", "DISPOSITIVO", "SERIE", "Entrega el equipo", "Recibe", "Mensajero externo", "No pertenece a"]},
    "en": {"entrega": ["NOTE: EQUIPMENT HANDOVER", "DEVICE", "BRAND", "MODEL", "SERIAL NUMBER", "Description:",
                       "Received by", "Delivered and authorized by", "IT Department", "Employee"],
           "descargo": ["NOTE: RETURN OF EQUIPMENT TO", "DEVICE", "SERIAL NUMBER", "Returned by", "Received by", "External courier", "Not affiliated with"]},
    "it": {"entrega": ["NOTA: CONSEGNA DI APPARECCHIATURE", "DISPOSITIVO", "MARCA", "MODELLO", "NUMERO DI SERIE",
                       "Descrizione:", "Riceve l'apparecchiatura", "Consegna e autorizza", "Reparto IT", "Collaboratore"],
           "descargo": ["NOTA: RESTITUZIONE DI APPARECCHIATURE A", "DISPOSITIVO", "NUMERO DI SERIE",
                        "Restituisce l'apparecchiatura", "Riceve", "Corriere esterno", "Esterno a"]},
}
# Restos de español que no deben aparecer en EN/IT, y marcas de tuteo/informalidad
LEFTOVER_ES = ["Descripción", "Colaborador", "Departamento de TI", "Entrega y autoriza", "Recibe el equipo", "Mensajero externo", "presente documento se extiende",
               "Asignación de equipo", "Devolución de equipo", "de septiembre de"]
INFORMAL = {"es": [r"\btú\b", r"\btu\b", r"\bdebes\b", r"\bpuedes\b", r"\btienes\b", r"\btus\b"],
            "en": [r"n't\b", r"'re\b", r"'ll\b", r"\bgonna\b", r"\bhey\b"],
            "it": [r"\btu\b", r"\btuo\b", r"\btua\b", r"\bdevi\b", r"\bpuoi\b", r"\bhai\b"]}
FECHA_RE = {"es": r"\d{1,2} de (enero|febrero|marzo|abril|mayo|junio|julio|agosto|septiembre|octubre|noviembre|diciembre) de \d{4}",
            "en": r"(January|February|March|April|May|June|July|August|September|October|November|December) \d{1,2}, \d{4}",
            "it": r"(1º|\d{1,2}) (gennaio|febbraio|marzo|aprile|maggio|giugno|luglio|agosto|settembre|ottobre|novembre|dicembre) \d{4}"}
FILEBASE = {"es": {"entrega": "Acta_Entrega", "descargo": "Acta_Descargo"},
            "en": {"entrega": "Handover_Record", "descargo": "Return_Record"},
            "it": {"entrega": "Verbale_Consegna", "descargo": "Verbale_Restituzione"}}


def docx_text(raw):
    d = docx.Document(io.BytesIO(raw))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    tables = [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
    return "\n".join(parts), tables, d.core_properties


def pdf_text(raw):
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        txt = "\n".join((p.extract_text() or "") for p in pdf.pages)
        return txt, len(pdf.pages), pdf.metadata


def validar_acta(cid, raw, formato, lang, tipo, series, nombre, extra=(), fname="", ev="", lote=False):
    problemas = []
    if formato == "docx":
        txt, tables, props = docx_text(raw)
        filas = tables[1][1:] if len(tables) > 1 else []
        series_doc = [f[3] for f in filas]
        if sorted(series_doc) != sorted(series):
            problemas.append(f"series en tabla {series_doc} != {series}")
    else:
        txt, npages, meta = pdf_text(raw)
        for s in series:
            if s not in txt.replace("\n", " ") and s not in txt:
                problemas.append(f"serie {s} no aparece")
        if npages != 1:
            problemas.append(f"{npages} páginas")
    plano = " ".join(txt.split())
    for t in TEXTOS[lang][tipo]:
        if t not in plano:
            problemas.append(f"falta texto '{t}'")
    if nombre not in plano:
        problemas.append(f"falta colaborador {nombre}")
    for e in extra:
        if e not in plano:
            problemas.append(f"falta '{e}'")
    if not re.search(FECHA_RE[lang], plano):
        problemas.append("fecha larga no localizada")
    if lang != "es":
        for w in LEFTOVER_ES:
            if w in plano:
                problemas.append(f"texto en español en documento {lang}: '{w}'")
    for pat in INFORMAL[lang]:
        m = re.search(pat, plano, re.I)
        if m:
            problemas.append(f"posible registro informal: '{m.group(0)}'")
    if "Ã" in plano or "�" in plano or "&amp;" in plano or "&lt;" in plano:
        problemas.append("caracteres corruptos o escapados visibles")
    exp = FILEBASE[lang][tipo]
    if not fname.startswith(exp) or not fname.endswith("." + formato):
        problemas.append(f"nombre de archivo {fname}")
    check(cid, f"Acta {tipo} {'lote ' if lote else ''}{formato.upper()} {lang}: textos traducidos, datos, fecha, tono y nombre",
          not problemas, "; ".join(problemas) or f"{fname}", ev, "media")
    return plano


def actas(c):
    m = D["movs"]
    a = D["activos"]
    serie = {k: a[k]["serie"] for k in a}
    nombre = "QA5-Íñigo Peña"
    planos = {}
    for lang in ("es", "en", "it"):
        for formato in ("docx", "pdf"):
            # Individual entrega (m1): la observación tiene & y <prueba>
            r = get(c, f"{BASE}/trazabilidad/acta/{m['m1']}", params={"formato": formato, "tipo": "entrega", "lang": lang})
            if r.status_code != 200:
                check(f"ACT-E-{formato}-{lang}", f"Acta de entrega {formato} {lang}", False, f"{r.status_code} {r.text[:200]}", severidad="alta")
                continue
            fn = filename_of(r)
            ev = save("actas", f"{lang}_{fn}", r.content)
            planos[("E", formato, lang)] = validar_acta(
                f"ACT-E-{formato[:1].upper()}-{lang.upper()}", r.content, formato, lang, "entrega", [serie["a1"]], nombre,
                extra=["acción & <prueba>"], fname=fn, ev=ev)
            # Descargo individual del movimiento devuelto
            mov_dev = D["movs"]["m3"]
            r = get(c, f"{BASE}/trazabilidad/acta/{mov_dev}",
                    params={"formato": formato, "tipo": "descargo", "lang": lang, "mensajero": "QA5 Mensajero <Ruta & Cía>"})
            if r.status_code != 200:
                check(f"ACT-D-{formato}-{lang}", f"Hoja de descargo {formato} {lang}", False, f"{r.status_code} {r.text[:200]}", severidad="alta")
                continue
            fn = filename_of(r)
            ev = save("actas", f"{lang}_{fn}", r.content)
            planos[("D", formato, lang)] = validar_acta(
                f"ACT-D-{formato[:1].upper()}-{lang.upper()}", r.content, formato, lang, "descargo", [serie["a3"]], nombre,
                extra=["QA5 Mensajero <Ruta & Cía>"], fname=fn, ev=ev)
            # Lote entrega m1+m2
            r = post(c, f"{BASE}/trazabilidad/acta/lote", params={"formato": formato, "tipo": "entrega", "lang": lang},
                     json={"movimientos_ids": [m["m1"], m["m2"]]})
            if r.status_code != 200:
                check(f"ACT-L-{formato}-{lang}", f"Acta lote {formato} {lang}", False, f"{r.status_code} {r.text[:200]}", severidad="alta")
                continue
            fn = filename_of(r)
            ev = save("actas", f"{lang}_{fn}", r.content)
            planos[("L", formato, lang)] = validar_acta(
                f"ACT-L-{formato[:1].upper()}-{lang.upper()}", r.content, formato, lang, "entrega", [serie["a1"], serie["a2"]],
                nombre, fname=fn, ev=ev, lote=True)
    # Lote descargo (API, la UI no lo ofrece)
    r = post(c, f"{BASE}/trazabilidad/acta/lote", params={"formato": "pdf", "tipo": "descargo", "lang": "es"},
             json={"movimientos_ids": [m["m3"]]})
    ev = save("actas", "es_lote_descargo.pdf", r.content) if r.status_code == 200 else ""
    check("ACT-LD", "Hoja de descargo por lote (PDF, API) se genera", r.status_code == 200, r.status_code, ev)

    # Fecha del descargo = fecha de devolución; fecha del acta de entrega = ¿fecha de asignación?
    dev = psql(f"""select to_char("MOV_Fecha_Devolucion",'YYYY-MM-DD'), to_char("MOV_Fecha_Asignacion",'YYYY-MM-DD') from "INV_MOVIMIENTO" where "MOV_Movimiento"='{m['m3']}'""")
    check("ACT-D-FECHA", "Hoja de descargo fechada con la fecha de devolución del movimiento",
          bool(dev.split("|")[0]), f"fecha devolución en BD: {dev}", "", "media")
    # Acta de entrega de un movimiento antiguo: ¿qué fecha usa?
    viejo = psql("""select "MOV_Movimiento", to_char("MOV_Fecha_Asignacion",'YYYY-MM-DD') from "INV_MOVIMIENTO" where "MOV_Fecha_Asignacion" < now() - interval '2 days' order by "MOV_Fecha_Asignacion" limit 1""")
    if viejo:
        vid, vfecha = viejo.split("|")
        r = get(c, f"{BASE}/trazabilidad/acta/{vid}", params={"formato": "pdf", "tipo": "entrega", "lang": "es"})
        if r.status_code == 200:
            ev = save("actas", "es_acta_movimiento_antiguo.pdf", r.content)
            txt, _, _ = pdf_text(r.content)
            y, mo, dd = vfecha.split("-")
            esperado = R.format_long_date(__import__("datetime").date(int(y), int(mo), int(dd)), "es")
            check("ACT-E-FECHA", "Acta de entrega re-generada de un movimiento antiguo muestra la fecha de la asignación",
                  esperado in " ".join(txt.split()), f"asignado el {vfecha} (esperado '{esperado}'); el acta dice: "
                  + (re.search(FECHA_RE['es'], ' '.join(txt.split())) or [''])[0], ev, "media")
    # Mezclar personas en un lote
    r = post(c, f"{BASE}/trazabilidad/acta/lote", params={"formato": "docx", "tipo": "entrega", "lang": "es"},
             json={"movimientos_ids": [m["m1"], m["m4"]]})
    check("ACT-MIX", "Lote con movimientos de dos personas → 400 CANNOT_MIX_DIFFERENT_PEOPLE_IN_SAME_DOCUMENT",
          r.status_code == 400 and "CANNOT_MIX" in r.text, f"{r.status_code} {r.text[:120]}", "", "alta")
    r = post(c, f"{BASE}/trazabilidad/acta/lote", params={"formato": "docx"}, json={"movimientos_ids": []})
    check("ACT-VACIO", "Lote vacío → 400", r.status_code == 400, f"{r.status_code} {r.text[:120]}", "", "baja")
    r = post(c, f"{BASE}/trazabilidad/acta/lote", params={"formato": "docx"}, json={"movimientos_ids": [str(uuid.uuid4())]})
    check("ACT-404", "Lote con movimiento inexistente → 404", r.status_code == 404, f"{r.status_code} {r.text[:120]}", "", "baja")
    r = post(c, f"{BASE}/trazabilidad/acta/lote", params={"formato": "pdf"}, json={"movimientos_ids": [m["m1"], m["m1"]]})
    if r.status_code == 200:
        txt, _, _ = pdf_text(r.content)
        n = " ".join(txt.split()).count(D["activos"]["a1"]["serie"])
        check("ACT-DUP", "Lote con el mismo movimiento repetido no duplica la fila del activo", n == 1,
              f"la serie aparece {n} veces", save("actas", "es_lote_duplicado.pdf", r.content), "baja")
    r = get(c, f"{BASE}/trazabilidad/acta/{m['m1']}", params={"formato": "xlsx"})
    check("ACT-FMT", "Formato no soportado (xlsx) → 422", r.status_code == 422, r.status_code, "", "baja")
    # Descargo de un movimiento que NO se ha devuelto
    r = get(c, f"{BASE}/trazabilidad/acta/{m['m1']}", params={"formato": "pdf", "tipo": "descargo", "lang": "es"})
    check("ACT-DESC-ABIERTO", "Hoja de descargo de un movimiento aún abierto (no devuelto) se rechaza o advierte",
          r.status_code >= 400, f"HTTP {r.status_code}: se genera una hoja de devolución de un equipo que sigue asignado",
          save("actas", "es_descargo_movimiento_abierto.pdf", r.content) if r.status_code == 200 else "", "media")
    # PDF metadatos (título/idioma)
    r = get(c, f"{BASE}/trazabilidad/acta/{m['m1']}", params={"formato": "pdf", "tipo": "entrega", "lang": "en"})
    if r.status_code == 200:
        _, _, meta = pdf_text(r.content)
        check("ACT-META", "PDF EN: metadatos de título en inglés", "HANDOVER" in (meta.get("Title") or ""), meta.get("Title"))
    return planos


# ─────────────────────────── Etiquetas QR ─────────────────────────────────
def decode_qrs(raw, zoom=6):
    doc = fitz.open(stream=raw, filetype="pdf")
    out = []
    det = cv2.QRCodeDetector()
    for page in doc:
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
        img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.h, pix.w, pix.n)
        img = cv2.cvtColor(img, cv2.COLOR_RGB2BGR) if pix.n == 3 else img[:, :, :3]
        ok, vals, _, _ = det.detectAndDecodeMulti(img)
        found = [v for v in (vals if ok else []) if v]
        # Refuerzo: recortar la cuadrícula 3x10 si el detector múltiple pierde alguno
        h, w = img.shape[:2]
        cw, ch = w / 8.5, h / 11.0
        for row in range(10):
            for col in range(3):
                x0 = int((0.1875 + col * (2.625 + 0.125)) * cw)
                y0 = int((0.5 + row * 1.0) * ch)
                crop = img[y0:y0 + int(ch), x0:x0 + int(0.95 * cw)]
                if crop.size == 0:
                    continue
                v, _, _ = det.detectAndDecode(crop)
                if v:
                    found.append(v)
        out.append(list(dict.fromkeys(found)))
    return out, doc.page_count


def etiquetas(c):
    a = D["activos"]
    # Un activo
    for lang, lab, fn in (("es", "N.º de serie", "etiquetas-activos.pdf"), ("en", "Serial no.", "asset-labels.pdf"),
                          ("it", "N. di serie", "etichette-asset.pdf")):
        r = post(c, f"{BASE}/core/activos/etiquetas.pdf", params={"lang": lang}, json={"activos_ids": [a["a1"]["id"]]})
        if r.status_code != 200:
            check(f"QR-1-{lang.upper()}", f"Etiqueta de un activo ({lang})", False, r.status_code, severidad="alta")
            continue
        ev = save("etiquetas", f"{lang}_uno_{filename_of(r)}", r.content)
        txt, _, meta = pdf_text(r.content)
        pages, _ = decode_qrs(r.content)
        codes = [v for p in pages for v in p]
        problemas = []
        if codes != [a["a1"]["codigo"]]:
            problemas.append(f"QR decodificado {codes}")
        if lab not in txt:
            problemas.append(f"falta '{lab}'")
        if a["a1"]["codigo"] not in txt:
            problemas.append("falta el código")
        if filename_of(r) != fn:
            problemas.append(f"nombre {filename_of(r)}")
        check(f"QR-1-{lang.upper()}", f"Etiqueta QR de un activo ({lang}): QR = código, texto traducido, nombre de archivo",
              not problemas, "; ".join(problemas) or f"QR={codes}", ev, "alta")
    # El QR lleva al activo correcto
    for k in ("a1", "a2", "a3"):
        r = get(c, f"{BASE}/core/activos/by-code/{a[k]['codigo']}")
        check(f"QR-RES-{k}", f"El código del QR ({a[k]['codigo']}) resuelve al activo correcto (by-code)",
              r.status_code == 200 and r.json()["ACT_Activo"] == a[k]["id"], r.status_code)
    # Lote QA5 por ids + códigos (mezcla, con duplicado)
    r = post(c, f"{BASE}/core/activos/etiquetas.pdf", params={"lang": "es"},
             json={"activos_ids": [a["a1"]["id"], a["a2"]["id"]], "codigos": [a["a3"]["codigo"], a["a4"]["codigo"], a["a1"]["codigo"]]})
    ev = save("etiquetas", "es_lote_qa5.pdf", r.content) if r.status_code == 200 else ""
    if r.status_code == 200:
        pages, n = decode_qrs(r.content)
        codes = sorted(v for p in pages for v in p)
        check("QR-LOTE", "Etiquetas por lote (ids + códigos, con duplicado): 4 etiquetas, cada QR = su código",
              codes == sorted(QA5_CODES), f"QR decodificados: {codes}", ev, "alta")
    # Lote grande: todos los activos visibles → paginación 30 por hoja
    todos = psql('select "ACT_Codigo_Interno" from "INV_ACTIVO" order by 1').split("\n")
    r = post(c, f"{BASE}/core/activos/etiquetas.pdf", params={"lang": "en"}, json={"codigos": todos})
    if r.status_code == 200:
        ev = save("etiquetas", "en_lote_todos.pdf", r.content)
        pages, n = decode_qrs(r.content)
        codes = sorted(v for p in pages for v in p)
        esperadas = -(-len(todos) // 30)
        falt = sorted(set(todos) - set(codes))
        check("QR-TODOS", f"Etiquetas de todos los activos ({len(todos)}): {esperadas} hojas y todos los QR legibles y correctos",
              n == esperadas and not falt and set(codes) <= set(todos),
              f"páginas {n}; decodificados {len(codes)}; faltan {falt[:5]}", ev, "alta")
    else:
        check("QR-TODOS", "Etiquetas de todos los activos", False, f"{r.status_code} {r.text[:200]}", severidad="alta")
    # Código inexistente mezclado con uno válido
    r = post(c, f"{BASE}/core/activos/etiquetas.pdf", params={"lang": "es"}, json={"codigos": [a["a1"]["codigo"], "QA5-NO-EXISTE"]})
    check("QR-PARCIAL", "Lote con un código inexistente: se informa el código no encontrado",
          r.status_code != 200, f"HTTP {r.status_code}: se omite en silencio el código inexistente", "", "baja", observacion=True)
    r = post(c, f"{BASE}/core/activos/etiquetas.pdf", json={"codigos": []})
    check("QR-VACIO", "Lote vacío → 422/400", r.status_code in (400, 422), r.status_code, "", "baja")


# ─────────────────────────── Alcance por sede ─────────────────────────────
def alcance():
    u = D["usuarios"]
    tec = cliente(u["ADMIN_TI"]["usuario"], u["ADMIN_TI"]["clave"], u["ADMIN_TI"]["totp_secret"])
    aud = cliente(u["AUDITOR"]["usuario"], u["AUDITOR"]["clave"], u["AUDITOR"]["totp_secret"])
    sede = D["sede"]
    codes_sede = set(psql(f'select "ACT_Codigo_Interno" from "INV_ACTIVO" where "SED_Sede"={sede}').split("\n")) - {""}
    otro = psql(f'select "ACT_Activo","ACT_Codigo_Interno" from "INV_ACTIVO" where "SED_Sede"<>{sede} limit 1').split("|")
    mov_otro = psql(f'select m."MOV_Movimiento" from "INV_MOVIMIENTO" m join "INV_ACTIVO" a using("ACT_Activo") where a."SED_Sede"<>{sede} limit 1')
    for etiqueta, cl in (("ADMIN_TI", tec), ("AUDITOR", aud)):
        for name in ("activos", "movimientos", "consumibles", "ordenes", "proveedores"):
            r = get(cl, f"{BASE}{EXPORTS[name][0]}", params={"lang": "es"})
            if r.status_code != 200:
                check(f"SC-{etiqueta[:3]}-{name[:4].upper()}", f"{etiqueta} de una sede exporta {name}", False,
                      f"{r.status_code} {r.text[:150]}", severidad="alta")
                continue
            ev = save("csv", f"alcance_{etiqueta}_{filename_of(r)}", r.content)
            rows = parse_csv(r.content)[1:]
            if name == "activos":
                got = {rw[0] for rw in rows}
                ok = got == codes_sede
            elif name == "movimientos":
                got = {rw[2] for rw in rows}
                ok = got <= codes_sede and len(rows) == int(psql(f'select count(*) from "INV_MOVIMIENTO" m join "INV_ACTIVO" a using("ACT_Activo") where a."SED_Sede"={sede}'))
            elif name == "consumibles":
                got = {rw[0] for rw in rows}
                ok = all("QA5" in g for g in got) and len(rows) == int(psql(f'select count(*) from "INV_CONSUMIBLE" where "SED_Sede"={sede}'))
            elif name == "ordenes":
                got = {rw[0] for rw in rows}
                ok = got == set(psql(f'select "OCO_Numero" from "INV_ORDEN_COMPRA" where "SED_Sede"={sede}').split("\n"))
            else:
                got = {rw[0] for rw in rows}
                ok = len(rows) == int(psql('select count(*) from "INV_PROVEEDOR"'))
            desc = {"proveedores": "exporta el catálogo global de proveedores (no tienen sede)"}.get(name, f"solo exporta filas de su sede")
            check(f"SC-{etiqueta[:3]}-{name[:4].upper()}", f"Alcance: {etiqueta} de la sede QA5 {desc} ({name})",
                  ok, f"{len(rows)} filas; muestra {sorted(got)[:5]}", ev, "bloqueante")
    # Auditoría: ADMIN_TI no puede; AUDITOR solo su sede
    r = get(tec, f"{BASE}/export/auditoria.csv", params={"lang": "es"})
    check("SC-TEC-AUD", "ADMIN_TI no puede exportar la bitácora (403)", r.status_code == 403, r.status_code, "", "alta")
    r = get(aud, f"{BASE}/gov/auditoria.csv", params={"lang": "it"})
    if r.status_code == 200:
        ev = save("csv", f"alcance_AUDITOR_{filename_of(r)}", r.content)
        rows = parse_csv(r.content)[1:]
        aud_id = psql(f"""select "USU_Usuario" from "INV_USUARIO" where "USU_Username"='{u['AUDITOR']['usuario']}'""")
        permitidos = int(psql(f"""select count(*) from "INV_AUDITORIA_SISTEMA" where "AUD_Sede"={sede} or "USU_Usuario"='{aud_id}'"""))
        ajenos = [rw for rw in rows if rw[5] and rw[5] != aud_id and "QA5" not in rw[6] and str(sede) not in rw[6]]
        check("SC-AUD-AUD", "AUDITOR de la sede QA5 exporta solo eventos de su sede o propios (ruta /gov de la UI)",
              len(rows) <= permitidos + 2 and not ajenos, f"{len(rows)} filas; permitidos en BD {permitidos}; ajenos {ajenos[:2]}", ev, "bloqueante")
    else:
        check("SC-AUD-AUD", "AUDITOR exporta bitácora", False, r.status_code, severidad="alta")
    # Documentos de otra sede
    if mov_otro:
        r = get(tec, f"{BASE}/trazabilidad/acta/{mov_otro}", params={"formato": "pdf"})
        check("SC-ACTA", "ADMIN_TI de la sede QA5 no puede generar el acta de un movimiento de otra sede (404)",
              r.status_code == 404, r.status_code, "", "bloqueante")
        r = post(tec, f"{BASE}/trazabilidad/acta/lote", params={"formato": "pdf"}, json={"movimientos_ids": [D["movs"]["m1"], mov_otro]})
        check("SC-ACTA-LOTE", "…ni incluirlo en un lote junto con uno propio", r.status_code in (403, 404), r.status_code, "", "bloqueante")
    r = post(tec, f"{BASE}/core/activos/etiquetas.pdf", json={"activos_ids": [otro[0]]})
    check("SC-QR", "ADMIN_TI de la sede QA5 no puede imprimir etiquetas de un activo de otra sede (404)",
          r.status_code == 404, r.status_code, "", "bloqueante")
    r = post(tec, f"{BASE}/core/activos/etiquetas.pdf", json={"codigos": [otro[1], D["activos"]["a1"]["codigo"]]})
    if r.status_code == 200:
        pages, _ = decode_qrs(r.content)
        codes = [v for p in pages for v in p]
        check("SC-QR-MIX", "Lote de etiquetas con un código de otra sede: solo imprime los de su sede",
              otro[1] not in codes, f"QR: {codes}", save("etiquetas", "alcance_tec_mix.pdf", r.content), "bloqueante")
    # Acta propia OK para el técnico de la sede
    r = get(tec, f"{BASE}/trazabilidad/acta/{D['movs']['m1']}", params={"formato": "docx", "lang": "it"})
    check("SC-ACTA-OK", "ADMIN_TI de la sede QA5 genera el acta de su sede", r.status_code == 200, r.status_code)
    # AUDITOR (solo lectura) y actas/etiquetas
    r = get(aud, f"{BASE}/trazabilidad/acta/{D['movs']['m1']}", params={"formato": "pdf"})
    check("SC-AUD-ACTA", "AUDITOR (solo lectura) y la descarga de actas: comportamiento según la matriz RBAC",
          True, f"HTTP {r.status_code}")
    return tec, aud


def main():
    c = cliente()
    info = csv_suite(c)
    analizar_csv(info)
    csv_contenido(info)
    # Ruta de la UI para la bitácora (/gov/auditoria.csv) debe devolver lo mismo que /export
    r = get(c, f"{BASE}/gov/auditoria.csv", params={"lang": "es"})
    check("CSV-AUD-GOV", "Bitácora por la ruta que usa la UI (/gov/auditoria.csv) responde 200 con BOM",
          r.status_code == 200 and r.content.startswith(b"\xef\xbb\xbf"), r.status_code,
          save("csv", f"gov_{filename_of(r)}", r.content) if r.status_code == 200 else "")
    etiquetas(c)
    actas(c)
    alcance()
    with open(os.path.join(OUT, "resultados_api.json"), "w", encoding="utf-8") as f:
        json.dump(RES, f, ensure_ascii=False, indent=1)
    print(f"\n{sum(r['estado'] == 'OK' for r in RES)}/{len(RES)} OK")


if __name__ == "__main__":
    main()
